"""Build the fictional fixture folder used by the browser check.

    pip install reportlab python-docx
    python3 experiments/studio/tests/make_fixtures.py

Everything here is invented ("Jordan Rivera"). Output: tests/fixtures/jordan-career/
"""
from pathlib import Path

root = Path(__file__).parent / "fixtures" / "jordan-career"
(root / "resumes" / "2022").mkdir(parents=True, exist_ok=True)
(root / "resumes" / "2024").mkdir(parents=True, exist_ok=True)
(root / "writing").mkdir(parents=True, exist_ok=True)

(root / "resumes" / "2022" / "resume.md").write_text(
    "# Jordan Rivera — résumé 2022\n\n## Experience\n\n"
    "- Reduced checkout latency by 35% by moving pricing to an edge cache.\n"
    "- Mentored four junior engineers through their first on-call rotation.\n"
    "- Led the migration of 40 services to the new deploy pipeline,\n  with zero customer-facing downtime.\n\n"
    "## Skills\n\nGo, TypeScript, Postgres 🐘\n", encoding="utf-8")
(root / "resumes" / "2024" / "resume.md").write_text(
    "# Jordan Rivera — résumé 2024\n\n## Experience\n\n"
    "- Mentored four junior engineers through their first on-call rotation.\n"
    "- Designed the incident review process now used by six teams.\n"
    "- Cut cloud spend 22% by right-sizing batch workers.\n", encoding="utf-8")
(root / "writing" / "brag-doc.txt").write_text(
    "Q3 wins\n\nShipped the accessibility audit fixes before the deadline.\n\n"
    "Wrote the onboarding guide that new hires now read first.\n", encoding="utf-8")
(root / "writing" / "photo.png").write_bytes(b"\x89PNG\r\n\x1a\nnot really")

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
c = canvas.Canvas(str(root / "resumes" / "jordan-portfolio.pdf"), pagesize=letter)
c.setFont("Helvetica-Bold", 16); c.drawString(72, 720, "Jordan Rivera - Portfolio")
c.setFont("Helvetica", 11)
c.drawString(72, 690, "Built a feature-flag service that handles 2 million evaluations per minute.")
c.drawString(72, 676, "It replaced three homegrown systems.")
c.drawString(72, 640, "Ran the design review for the payments ledger rewrite.")
c.showPage()
c.setFont("Helvetica", 11)
c.drawString(72, 720, "Page two: volunteered as a coding mentor at the public library.")
c.save()

from docx import Document
d = Document()
d.add_heading("Cover letter draft", level=1)
d.add_paragraph("I care about calm, observable systems and teams that learn from incidents.")
d.add_paragraph("Introduced blameless postmortems across the platform group.", style="List Bullet")
d.add_paragraph("Partnered with support to halve repeat tickets.", style="List Bullet")
d.save(str(root / "writing" / "cover-letter.docx"))
print("wrote", root)
